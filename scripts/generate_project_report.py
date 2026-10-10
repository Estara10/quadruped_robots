#!/usr/bin/env python3
"""Generate a comprehensive ABS project report (.docx) for research reporting."""

import hashlib
import subprocess
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path.home() / "quadruped_robots"
MODELS = {
    "agile_policy": ROOT / "quadruped_ros2_control_humble" / "descriptions" / "unitree" / "go2_description" / "config" / "abs" / "policy.pt",
    "recovery_policy": ROOT / "quadruped_ros2_control_humble" / "descriptions" / "unitree" / "go2_description" / "config" / "rec" / "policy.pt",
    "ra_value": ROOT / "quadruped_ros2_control_humble" / "descriptions" / "unitree" / "go2_description" / "config" / "abs" / "ra_value.pt",
}


def git_log() -> str:
    try:
        return subprocess.check_output(
            ["git", "log", "--oneline", "-20"],
            cwd=ROOT, text=True,
        ).strip()
    except Exception:
        return "(git unavailable)"


def model_hashes() -> dict:
    hashes = {}
    for name, path in MODELS.items():
        if path.exists():
            hashes[name] = hashlib.sha256(path.read_bytes()).hexdigest()[:16] + "..."
        else:
            hashes[name] = "MISSING"
    return hashes


def build() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = "Calibri"

    # ── Title ──
    title = doc.add_heading("ABS 论文复现项目 — 工作总结报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # ── 1. 项目目标 ──
    doc.add_heading("1. 项目目标", level=1)
    doc.add_paragraph(
        "复现 ABS (Agile But Safe) 论文在 Unitree Go2 四足机器人上的完整控制链路："
        "敏捷策略 + RA 值网络 + 恢复策略 + 目标导航 + 深度射线感知。"
        "仿真平台为 MuJoCo + ROS2 Humble + LibTorch，实机为 Go2 + ZED 相机。"
    )

    # ── 2. 系统架构 ──
    doc.add_heading("2. 系统架构", level=1)
    doc.add_paragraph(
        "仿真管线: MuJoCo (物理 + 射线) → DDS → ros2_control → StateRL.cpp (策略推理 / RA / recovery)\n"
        "实机管线: Go2 DDS → ros2_control → StateRL.cpp (同一套代码)\n"
        "感知: 仿真用 MuJoCo 几何 ray2d; 实机用 ZED 深度 → ResNet18 → 11 条 log2 射线"
    )

    # ── 3. 已完成工作 ──
    doc.add_heading("3. 已完成工作", level=1)

    doc.add_heading("3.1 控制算法复现", level=2)
    items = [
        "敏捷策略推理 (61-dim 观测 → 12 关节目标, TorchScript)",
        "RA 值网络 (19→64→64→1 Tanh), 按阈值触发 recovery",
        "恢复策略推理 (49-dim 观测, 触发期间内联替代 agile action)",
        "Recovery Twist 梯度下降优化 (论文方法: 3 次迭代, loss=λ·max(ra+2ε,0)+0.02·pos²)",
        "世界坐标目标导航 + 到达检测 + 停止",
        "FSM 状态机: PASSIVE → FIXEDDOWN → FIXEDSTAND → RL",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3.2 仿真平台搭建", level=2)
    items = [
        "MuJoCo Go2 模型 + 多场景 (平坦/障碍/地形/坡地/随机障碍)",
        "DDS 桥: MuJoCo ↔ ros2_control (LowCmd / LowState / SportModeState)",
        "几何 ray2d: 2D 射线圆/矩形相交, 11 条射线 (-45°~+45°), log2 输出",
        "碰撞检测: /mujoco_collision 共享内存, 机器人 vs 障碍物 contact 统计",
        "结构化评估: 自动启动 → 监控 → 统计 (run_abs_eval.py + analyze_abs_eval.py)",
        "仿真基线: 4 场景 × 3 次 = 12/12 全部成功 (100%), 零跌倒",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3.3 Ray-Pred 深度射线感知", level=2)
    items = [
        "离线对比工具: compare_ray_pred.py (ResNet18 预测 vs 几何真值, 不写控制 shm)",
        "发现并修复 MuJoCo 深度相机参数错误 (朝向反了, 位置/FOV 未对齐训练配置)",
        "MuJoCo 数据集生成 + 多轮微调 (数据增强, 加权 loss, 安全评分)",
        "MUJOCO_RAY_SOURCE 开关: geometric (默认) / ray_pred (ResNet18)",
        "一键 Ray-Pred 仿真启动脚本",
        "ZED 实机射线预测器 (zed_ray_predictor.py)",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3.4 安全机制", level=2)
    items = [
        "姿态安全: roll/pitch > 75° → PASSIVE (IMU 四元数)",
        "DDS 超时: 关节位置 100 步无变化 → PASSIVE",
        "关节限位: hip ±1.047, thigh [-1.57, 3.49], calf [-2.72, -0.84]",
        "Action 输出限幅: ±4.0 (≈ ±1 rad 关节偏移)",
        "软启动: Kp/Kd 从 0 渐变到目标值 (250 步)",
        "遥控器急停: command=1 → PASSIVE (匹配原文 B 键)",
        "扭矩监测: PD 扭矩 > 额定 × 2.5 → PASSIVE (匹配原文 PowerProtect)",
    ]
    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    # ── 4. 模型与配置 ──
    doc.add_heading("4. 关键模型与配置", level=1)
    hashes = model_hashes()
    for name, h in hashes.items():
        doc.add_paragraph(f"{name}: SHA256 {h}")

    doc.add_paragraph()
    doc.add_paragraph("配置参数:")
    config_items = [
        "goal_x=7.0m, goal_y=0.0m, resample_goal_on_arrival=false (到达即停)",
        "Kp=30, Kd=0.65, action_scale=0.25, decimation=4",
        "RA threshold=-0.05 (进入) / -0.08 (退出)",
        "recovery_hold_steps=30, twist_lam=10, twist_lr=0.5",
        "contact_threshold=1N (足力, 匹配训练)",
        "policy_joint_order=ros1_fl_fr_rl_rr (策略关节顺序 remap)",
    ]
    for item in config_items:
        doc.add_paragraph(item, style="List Bullet")

    # ── 5. 当前状态 ──
    doc.add_heading("5. 当前状态与下一步", level=1)
    doc.add_paragraph(
        "仿真: 核心链路端到端验证通过, 几何 ray2d 条件下 100% 成功率。"
        "Ray-Pred 仿真可用但非必须——实机天然匹配 ZED 训练分布, 无需 MuJoCo 域适应。"
    )
    doc.add_paragraph(
        "代码分支: feat/ray-pred-source-switch (5 commits)"
    )
    doc.add_paragraph(
        "下一步: 实机 Go2 部署。笔记本电脑 → 连 Go2 网线 → 改 network_interface → "
        "首测 goal_x=1.5m 低速短距 → 逐步增加距离。"
    )

    # ── 6. 提交历史 ──
    doc.add_heading("6. 提交历史", level=1)
    for line in git_log().split("\n"):
        doc.add_paragraph(line, style="No Spacing")

    # ── 7. 技术说明 ──
    doc.add_heading("7. 关键实现细节", level=1)
    doc.add_paragraph(
        "• 关节顺序: 全链路 FR, FL, RR, RL (MuJoCo/DDS/Controller), "
        "策略为 FL-first, 通过 policy_joint_order=ros1_fl_fr_rl_rr remap\n"
        "• Timer 恒为 0.5 (匹配 ROS1 部署), Contact = +1(着地)/-1(离地)\n"
        "• RA/recovery 的 lin_vel 必须是机体系速度, MuJoCo 优先用 odometer world velocity 旋到 body frame\n"
        "• 频率: ROS2 125Hz vs ROS1 12.5Hz, recovery_hold_steps=30 匹配 ROS1 约 3 步有效时长\n"
        "• 安全机制独立于控制器状态, 任一触发即切 PASSIVE (归零 Kp/Kd)"
    )

    # ── 8. 文件清单 ──
    doc.add_heading("8. 主要文件清单", level=1)
    files = [
        "quadruped_ros2_control_humble/controllers/rl_quadruped_controller/src/FSM/StateRL.cpp",
        "unitree_mujoco/simulate/src/unitree_sdk2_bridge.h",
        "scripts/launch_abs_obstacle.sh (几何 ray 仿真)",
        "scripts/launch_abs_ray_pred.sh (Ray-Pred 仿真)",
        "scripts/launch_abs_real.sh (实机启动)",
        "scripts/zed_ray_predictor.py (实机 ZED 射线)",
        "scripts/compare_ray_pred.py (Ray-Pred 离线评估)",
        "scripts/finetune_mujoco_ray_pred.py (模型微调)",
        "scripts/generate_mujoco_ray_dataset.py (数据集生成)",
        "scripts/run_abs_eval.py (结构化仿真评估)",
        "scripts/analyze_abs_eval.py (评估报告生成)",
        "quadruped_ros2_control_humble/descriptions/unitree/go2_description/config/abs/config.yaml",
    ]
    for f in files:
        doc.add_paragraph(f, style="List Bullet")

    # Save
    output = ROOT / "ABS项目工作总结报告.docx"
    doc.save(str(output))
    return output


if __name__ == "__main__":
    path = build()
    print(f"Report saved: {path}")
